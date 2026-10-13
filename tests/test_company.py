"""Company Excel layouts, Express files and the code cross-reference."""

import pytest

from conftest import SAMPLES
from partlist_studio.core import stock as stock_mod
from partlist_studio.core.company_files import (SHEET_BOM, SHEET_PARTS_BALANCE, SHEET_STEEL, SHEET_STOCK, ModelBOM,
                                                ModelRow, export_models, export_stock_count, models_from_rows,
                                                models_to_rows, normalize_code, parse_qty, read_company_workbook,
                                                split_codes)
from partlist_studio.core.express import (check_part_codes, code_issues, compare_stock, export_new_items,
                                          export_pd_draft, pd_lines, read_express_file, thai_date)
from partlist_studio.core.models import Part, PartList
from partlist_studio.core.stock import StockItem, StockList, check_requirements, sheet_size_mm, sheet_weight_kg
from partlist_studio.core.tabular import read_table
from partlist_studio.core.xref import CodeXref, XrefEntry, suggest_links

NBSP = "\xa0"


@pytest.fixture
def company_book(tmp_path):
    """A small workbook in the layout of 'Parts List All Model' (with a picture) + stock sheets."""
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "กระบะคอกสูง3ราว"
    ws.append(["รายการชิ้นส่วนประกอบกระบะ Toyota (คอกสูง 3 ราว)"])
    ws.append(["ลำดับ", "รหัส", "รายการ", "ภาพประกอบ", "ขนาด", "จำนวนที่ใช้", "หมายเหตุ"])
    ws.append(["1", "RRS-A-107", "เฟรมล่าง  ", None, None, "1", "Lower frame"])
    ws.append(["2", "RRS-T018", "พื้น คอกสูง", None, None, "1+1", "ซ้าย-ขวา"])
    ws.append(["3", "RRS-T-028,029", "เหล็กซองเสา", None, None, "2+2+2", ""])
    ws.append(["4", "RRS-A-029-เสาหน้า , RRS-A-028-เสาหลัง", "เหล็กซองเสา", None, None, "2", ""])
    pic = tmp_path / "p.png"
    Image.new("RGB", (60, 40), "orange").save(pic)
    ws.add_image(XLImage(str(pic)), "D3")
    st = wb.create_sheet("สต๊อกเหล็ก")
    st.append(["ลำดับ", "รายการ", "ยอดยกมา", "ยอดรับเข้า", "จ่ายออก", None, "รวมจ่าย", "คงเหลือ"])
    st.append([None, None, None, None, "2025-07-01", "2025-07-02", None, None])
    st.append([1, "เหล็กแผ่นดำ 4 มม.  5x10 ฟุต", 14, 2, 1, None, 1, 15])
    st.append([2, "ท่อเหล็กกลม ดำ  1  นิ้ว", 78, None, None, None, 0, None])
    ap = wb.create_sheet("อะไหล่กระบะ")
    ap.append(["รายงานยอดคงเหลือ", None, None, "วันที่ 07/09/24"])
    ap.append(["ลำดับ", "รายการ", "ยอดคงเหลือ", "หมายเหตุ"])
    ap.append(["1", "หน้าโปร 40 (คอกเตี้ย)", "3 คัน"])
    ap.append(["2", "เสา 40", "68 อัน"])
    path = tmp_path / "parts_list_all_model.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def count_file(tmp_path):
    """Layout of '10. รายการ Stock ชิ้นส่วนกระบะดั้มพ์.นับสต๊อค' (NBSP in codes and names)."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "STBALLL"
    ws.append(["รหัสสินค้า", "รายการ", "จำนวน", "หน่วย", "หมายเหตุ"])
    ws.append([f"01-0001-0001{NBSP}", f"{NBSP}หน้าโปร{NBSP}80", 0, "ชุด", "รอตัด"])
    ws.append([f"01-0007-0004{NBSP}", f"{NBSP}2ราวข้าง(209)กระบะ2.50", 6, "ชิ้น", ""])
    form = wb.create_sheet("Format Stock R1")
    form.append(["Rungrueng group"])
    form.append(["Section :  RRS Dump"])
    form.append([])
    form.append([])
    form.append(["รหัสสินค้า", "รายการ", "Minimum (pcs.)", "ยอดคงเหลือ (pcs.)", "ยอดคงเหลือ", "หมายเหตุ"])
    form.append([f"01-0001-0001{NBSP}", f"{NBSP}หน้าโปร{NBSP}80", 2, 0, "ชุด", "รอตัด"])
    form.append([f"01-0007-0004{NBSP}", f"{NBSP}2ราวข้าง(209)กระบะ2.50", "", 6, "ชิ้น", ""])
    path = tmp_path / "stock_count.xlsx"
    wb.save(path)
    return path


# ---- codes / quantities ---------------------------------------------------------
@pytest.mark.parametrize("raw, codes", [
    ("RRS-T018", ["RRS-T-018"]),
    ("RRS-A_102", ["RRS-A-102"]),
    ("RRS-A--033", ["RRS-A-033"]),
    ("RRS-A035", ["RRS-A-035"]),
    ("RRS-T-028,029", ["RRS-T-028", "RRS-T-029"]),
    ("RRS-A-029-เสาหน้า , RRS-A-028-เสาหลัง", ["RRS-A-029", "RRS-A-028"]),
    ("RRS-A-029เสาหน้า,RRS-A-028เสาหลัง", ["RRS-A-029", "RRS-A-028"]),
    ("RRS-A-131(R) , RRS-A-131(L)", ["RRS-A-131(R)", "RRS-A-131(L)"]),
    ("RRS-M-009( R)", ["RRS-M-009(R)"]),
    ("", []),
])
def test_split_codes(raw, codes):
    assert split_codes(raw) == codes


def test_normalize_keeps_good_codes():
    assert normalize_code("RRS-A-108-209-2-50") == "RRS-A-108-209-2-50"
    assert normalize_code("DP-ST-RRS-066") == "DP-ST-RRS-066"


@pytest.mark.parametrize("text, qty", [("1+1", 2), ("2+2+2", 6), ("6", 6), ("", 1), ("10", 10)])
def test_parse_qty(text, qty):
    assert parse_qty(text)[0] == qty


# ---- company workbook --------------------------------------------------------------
def test_read_company_workbook(company_book, tmp_path):
    book = read_company_workbook(company_book, tmp_path / "img")
    assert book.kinds == {"กระบะคอกสูง3ราว": SHEET_BOM, "สต๊อกเหล็ก": SHEET_STEEL, "อะไหล่กระบะ": SHEET_PARTS_BALANCE}
    [model] = book.models
    assert model.title.startswith("รายการชิ้นส่วนประกอบกระบะ")
    assert [r.codes for r in model.rows] == [["RRS-A-107"], ["RRS-T-018"], ["RRS-T-028", "RRS-T-029"],
                                             ["RRS-A-029", "RRS-A-028"]]
    assert [r.qty for r in model.rows] == [1, 2, 6, 2]
    assert model.rows[0].image and model.rows[0].image.endswith(".jpg")  # picture anchored on its row
    steel = book.stocks["สต๊อกเหล็ก"]
    plate, pipe = steel.items
    assert (plate.on_hand, plate.unit, plate.spec) == (15, "แผ่น", "t4")
    assert pipe.on_hand == 78 and pipe.unit == "เส้น"  # คงเหลือ empty -> ยกมา + รับ - จ่าย
    assert book.steel_dates["สต๊อกเหล็ก"] == ["2025-07-01", "2025-07-02"]
    parts = book.stocks["อะไหล่กระบะ"]
    assert [(i.name, i.on_hand, i.unit) for i in parts.items] == [("หน้าโปร 40 (คอกเตี้ย)", 3, "คัน"),
                                                                  ("เสา 40", 68, "อัน")]


def test_model_to_parts_and_round_trip(company_book, tmp_path):
    model = read_company_workbook(company_book, tmp_path / "img").models[0]
    parts = model.to_parts(sets=2)
    assert parts[1].qty == 4 and "1+1" in parts[1].remark
    assert "RRS-T-029" in parts[2].remark  # second code kept
    out = tmp_path / "out.xlsx"
    export_models([model], out)
    again = read_company_workbook(out, tmp_path / "img2").models[0]
    assert [(r.codes, r.qty, r.name) for r in again.rows] == [(r.codes, r.qty, r.name) for r in model.rows]
    assert again.rows[0].image  # the picture is written back into ภาพประกอบ


def test_models_flat_rows_round_trip():
    models = [ModelBOM("A", "แบบ A", [ModelRow("1", "RRS-A-1", ["RRS-A-1"], "x", "", "1+1", 2, "", "pic.jpg")])]
    again = models_from_rows(models_to_rows(models), images_from=models)
    assert again[0].rows[0].qty == 2 and again[0].rows[0].image == "pic.jpg"


def test_stock_count_file_both_sheets(count_file, tmp_path):
    book = read_company_workbook(count_file)
    assert book.kinds == {"STBALLL": SHEET_STOCK, "Format Stock R1": SHEET_STOCK}
    for sheet in ("STBALLL", "Format Stock R1"):
        items = book.stocks[sheet].items
        assert [(i.code, i.name, i.on_hand, i.unit) for i in items] == [
            ("01-0001-0001", "หน้าโปร 80", 0, "ชุด"), ("01-0007-0004", "2ราวข้าง(209)กระบะ2.50", 6, "ชิ้น")]
    assert book.stocks["Format Stock R1"].items[0].min_qty == 2
    assert book.stocks["STBALLL"].items[0].remark == "รอตัด"
    out = tmp_path / "count_out.xlsx"
    export_stock_count(book.stocks["Format Stock R1"], out)
    again = read_company_workbook(out)
    assert {k: [(i.code, i.on_hand, i.unit) for i in v.items] for k, v in again.stocks.items()} == {
        "STBALLL": [("01-0001-0001", 0, "ชุด"), ("01-0007-0004", 6, "ชิ้น")],
        "Format Stock R1": [("01-0001-0001", 0, "ชุด"), ("01-0007-0004", 6, "ชิ้น")]}


# ---- Express -------------------------------------------------------------------
def test_express_text_report():
    stock, report = read_express_file(SAMPLES / "express_stock_count_report.txt")
    assert report.title == "แบบฟอร์มเพื่อการตรวจนับสินค้า"
    assert report.warehouse.startswith("01")
    codes = {i.code: i for i in stock.items}
    assert codes["RRS-A-001-220"].name == "ข้างเฟรมบน 220" and not codes["RRS-A-001-220"].code_truncated
    cut = codes["RRS-A-108-209-2-"]
    assert cut.code_truncated and cut.name == "2ราวข้าง(209)กระบะ2.50"
    assert codes["01-B00-0015"].unit == "ผน" and codes["01-B00-0015"].spec == "t4"
    assert not stock.has_quantities  # a count form has no balances


def test_express_item_list_and_truncated_lookup():
    full, _ = read_express_file(SAMPLES / "express_item_list.csv")
    item = next(i for i in full.items if i.code == "01-B00-0015")
    assert (item.category, item.account_group, item.unit_name, item.on_hand, item.unit_cost) == (
        "01", "RM01", "แผ่น", 6, 3226)
    report, _ = read_express_file(SAMPLES / "express_stock_count_report.txt")
    part = Part(part_no="RRS-A-108-209-2-50")
    assert stock_mod.find_by_code(part.part_no, report).code == "RRS-A-108-209-2-"
    rows = compare_stock(full, report, "Excel", "รายงาน")
    assert next(r for r in rows if r.code == "RRS-A-108-209-2-50").status == "ตรงกัน"


def test_sheet_conversion_and_pd_draft(tmp_path):
    full, _ = read_express_file(SAMPLES / "express_item_list.csv")
    plate = next(i for i in full.items if i.code == "01-B00-0015")
    assert sheet_size_mm(plate.name) == (1524, 3050)
    assert sheet_size_mm("เหล็กแผ่นดำ 3มม. 4x2530") == pytest.approx((1219.2, 2530))
    assert sheet_weight_kg(plate) == pytest.approx(1.524 * 3.05 * 4 * 7.85, rel=1e-3)
    pl = PartList(parts=[Part(part_no="P1", material="SS400", spec="t4 x 900 x 400", qty=2, weight=50)])
    pl.info.build_qty = 3
    lines = check_requirements(pl, full)
    line = next(l for l in lines if l.code == "01-B00-0015")
    assert line.required == pytest.approx(300 / sheet_weight_kg(plate))  # 300 kg of 4 mm plate in sheets
    assert "kg/แผ่น" in line.note
    rows = pd_lines(lines)
    assert rows[0] == [1, "01-B00-0015", plate.name, "01", 3, "ผน", 3226, 9678]  # whole sheets
    out = tmp_path / "pd.xlsx"
    assert export_pd_draft(pl, lines, out, department="ขาย1", doc_date="2026-10-10") == 1
    table = read_table(out)
    assert table[1][:2] == ["สำหรับแผนก", "ขาย1"] and "10/10/69" in table[1]


def test_new_items_and_code_checks(tmp_path):
    full, _ = read_express_file(SAMPLES / "express_item_list.csv")
    pl = PartList(parts=[Part(part_no="RRS-A-108-209-2-50", name="2ราวข้าง"), Part(part_no="RRS-X 1", name="ใหม่")])
    checks = check_part_codes(pl, full)
    assert [c.status for c in checks] == ["ตรงกัน", "ไม่พบใน Express"]
    assert "ช่องว่าง" in checks[1].detail
    assert export_new_items(pl, full, tmp_path / "new.xlsx") == 1
    assert thai_date("2026-10-10") == "10/10/69"
    assert code_issues("RRS-A-001") == []


# ---- code cross-reference -----------------------------------------------------------
def test_xref_links_express_and_count_codes(count_file):
    express, _ = read_express_file(SAMPLES / "express_item_list.csv")
    counted = read_company_workbook(count_file).stocks["STBALLL"]
    links = suggest_links(express, counted)
    assert [(e.code_a, e.code_b) for e in links] == [("RRS-A-108-209-2-50", "01-0007-0004")]
    xref = CodeXref(links)
    old = stock_mod.XREF_LOOKUP
    stock_mod.XREF_LOOKUP = xref.lookup
    try:
        found = stock_mod.find_by_code("RRS-A-108-209-2-50", counted)
        assert found.code == "01-0007-0004" and found.on_hand == 6
    finally:
        stock_mod.XREF_LOOKUP = old
    again = CodeXref.from_rows(xref.to_rows())
    assert again.entries[0].code_b == "01-0007-0004"
    assert xref.merge([XrefEntry("RRS-A-108-209-2-50", "01-9999")]) == 0  # never replaces a link


def test_unit_column_found_by_content():
    rows = [["รหัสสินค้า", "รายการ", "ยอดคงเหลือ (pcs.)", "ยอดคงเหลือ"], ["A", "x", 2, "ชิ้น"], ["B", "y", 3, "ต้น"]]
    items = StockList.from_rows(rows).items
    assert [(i.on_hand, i.unit) for i in items] == [(2, "ชิ้น"), (3, "ต้น")]
    assert StockItem(name="เหล็กแผ่นดำ 5มม.").is_sheet
