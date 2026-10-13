"""Item master workbook ("ฟอร์มทะเบียนสินค้า") built from Express STMASR0200.

Tabs: คำแนะนำ | ทะเบียนสินค้า | สรุป | กลุ่มรหัส | สต็อกนับ | เทียบรหัสแบบกระบะ | ฟอร์มเพิ่มสินค้าใหม่ |
รายการเลือก | ความหนาแน่น.  Imported data are typed values; everything derived (meaning of the code
group, density, weight per sheet, stock balance, code checks, summary counts) is a formula over those
cells, so the sheet stays right when someone edits it in Google Sheets.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from .company_files import ModelBOM, base_code
from .dwg_io import DEFAULT_DENSITIES, parse_thickness
from .express import CODE_FAMILY_NAMES, code_family, item_kind
from .stock import MATERIAL_FAMILIES, StockItem, StockList, _key, sheet_size_mm
from .xref import CodeXref

MASTER = "ทะเบียนสินค้า"
FAMILIES = "กลุ่มรหัส"
COUNTED = "สต็อกนับ"
BOMCHECK = "เทียบรหัสแบบกระบะ"
FORM = "ฟอร์มเพิ่มสินค้าใหม่"
LISTS = "รายการเลือก"
DENSITY = "ความหนาแน่น"
SUMMARY = "สรุป"
GUIDE = "คำแนะนำ"

MASTER_HEADERS = [
    "รหัสสินค้า", "ชื่อสินค้า", "หมวด", "ชื่อหมวด", "กลุ่มรหัส", "ความหมายกลุ่มรหัส", "ประเภท",
    "หน่วยย่อย", "หน่วยใหญ่", "ตัวคูณใหญ่", "หน่วยซื้อ", "ตัวคูณซื้อ", "หน่วยขาย", "ตัวคูณขาย",
    "รหัสบ/ช", "ราคามาตรฐาน", "สินค้าทดแทน", "ผู้จัดจำหน่าย", "หมายเหตุ (Express)",
    "กลุ่มวัสดุ", "ความหนา (มม.)", "กว้าง (มม.)", "ยาว (มม.)", "ความหนาแน่น (kg/dm³)", "น้ำหนัก/แผ่น (kg)",
    "รหัสนับสต็อก", "ยอดคงเหลือ (นับสต็อก)", "หมายเหตุนับสต็อก", "ตรวจรหัส",
]
FORM_HEADERS = ["รหัสสินค้า", "ชื่อสินค้า (ชื่อไทย)", "ชื่ออังกฤษ", "หมวด", "กลุ่มรหัส", "หน่วยย่อย", "หน่วยซื้อ",
                "ตัวคูณซื้อ", "หน่วยขาย", "ตัวคูณขาย", "รหัสบ/ช", "ราคามาตรฐาน", "ผู้จัดจำหน่าย", "หมายเหตุ",
                "ผู้ขอ", "วันที่ขอ", "ตรวจรหัส"]
UNITS = ["ชิ้น", "ชุด", "ต้น", "แผ่น", "เส้น", "ตัว", "กล่อง", "อัน", "คู่", "กก.", "ม."]
FAMILY_DENSITY = {"STEEL": "SS400", "SUS": "SUS", "AL": "AL", "GALV": "SGCC"}


def _family(item: StockItem) -> str:
    upper = (item.material or item.name).upper().replace(" ", "")
    for fam, words in MATERIAL_FAMILIES:
        if any(w.upper() in upper for w in words):
            return fam
    if any(w in item.name for w in ("เหล็ก", "ท่อ", "แฟลตบาร์", "ตะแกรง")):
        return "STEEL"
    return ""


def _density_rows(densities: dict[str, float]) -> list[list]:
    rows = [["STEEL", densities.get("SS400", 7.85), "เหล็กดำ / SS400 / SPHC / DC01"],
            ["SUS", densities.get("SUS", 7.93), "สแตนเลส"],
            ["AL", densities.get("AL", 2.70), "อลูมิเนียม"],
            ["GALV", densities.get("SGCC", 7.85), "เหล็กชุบสังกะสี / กัลวาไนซ์"]]
    return rows


def build_item_form(express: StockList, path: str | Path, counted: StockList | None = None,
                    models: list[ModelBOM] | None = None, xref: CodeXref | None = None,
                    densities: dict[str, float] | None = None, source: str = "STMASR0200",
                    resolved: dict | None = None, form_rows: list[list] | None = None,
                    balances: dict[str, StockList] | None = None) -> dict[str, int]:
    """With `resolved` (express_codes.resolve_models) the models are written as tabs in their own layout
    with Express codes - see partlist_express.build_partlist_express."""
    from types import SimpleNamespace

    import openpyxl
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    densities = densities or DEFAULT_DENSITIES
    counted = counted or StockList()
    models = models or []
    xref = xref or CodeXref()
    font = "Arial"
    head_fill = PatternFill("solid", fgColor="1F4E79")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    head_font = Font(name=font, bold=True, color="FFFFFF")
    body_font = Font(name=font, size=10)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def sheet(title: str, headers: list[str], widths: list[int] | None = None, header_row: int = 1):
        return header(wb.create_sheet(title), headers, widths, header_row)

    def header(ws, headers: list[str], widths: list[int] | None = None, header_row: int = 1):
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(header_row, c, h)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = ws.cell(header_row + 1, 1)
        for i, w in enumerate(widths or [], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[header_row].height = 32
        return ws

    def style_body(ws, first_row: int, last_row: int, cols: int):
        for row in ws.iter_rows(min_row=first_row, max_row=last_row, max_col=cols):
            for cell in row:
                cell.font = body_font

    guide = wb.create_sheet(GUIDE)

    # ---- master -----------------------------------------------------------------------
    ws = sheet(MASTER, MASTER_HEADERS, [20, 46, 6, 24, 9, 24, 11, 8, 8, 7, 8, 7, 8, 7, 8, 10, 12, 14, 20,
                                        9, 9, 9, 9, 10, 11, 14, 11, 16, 30])
    counted_by_code = {_key(i.code): i for i in counted.items}
    n = len(express.items)
    for r, item in enumerate(express.items, start=2):
        x = item.extra
        fam = _family(item)
        size = sheet_size_mm(item.name) if item.is_sheet else None
        thk = parse_thickness(item.spec) or None
        count_code = next((c for c in xref.lookup(item.code) if _key(c) in counted_by_code), "")
        values = [item.code, item.name, item.category, x.get("category_name", ""), code_family(item.code), None,
                  item_kind(item.code), item.unit, x.get("unit_big", ""), x.get("factor_big"), x.get("unit_buy", ""),
                  x.get("factor_buy"), x.get("unit_sell", ""), x.get("factor_sell"), item.account_group,
                  item.unit_cost, x.get("substitute", ""), x.get("vendor", ""), item.remark,
                  fam, thk, round(size[0], 1) if size else None, round(size[1], 1) if size else None]
        for c, v in enumerate(values, start=1):
            if v not in (None, ""):
                ws.cell(r, c, v)
        ws.cell(r, 6, f'=IFERROR(VLOOKUP(E{r},{FAMILIES}!$A:$B,2,FALSE),"")')
        ws.cell(r, 24, f'=IF(T{r}="","",IFERROR(VLOOKUP(T{r},{DENSITY}!$A:$B,2,FALSE),""))')
        ws.cell(r, 25, f'=IF(AND(ISNUMBER(U{r}),ISNUMBER(V{r}),ISNUMBER(W{r}),ISNUMBER(X{r})),'
                       f'ROUND(U{r}*V{r}*W{r}*X{r}/1000000,2),"")')
        if count_code:
            ws.cell(r, 26, count_code)
        ws.cell(r, 26).fill = input_fill
        ws.cell(r, 27, f'=IF(Z{r}="","",IFERROR(INDEX({COUNTED}!$C:$C,MATCH(Z{r},{COUNTED}!$A:$A,0)),'
                       f'"ไม่พบรหัสนี้"))')
        ws.cell(r, 28, f'=IF(Z{r}="","",IFERROR(INDEX({COUNTED}!$E:$E,MATCH(Z{r},{COUNTED}!$A:$A,0))&"",""))')
        ws.cell(r, 29, f'=TRIM(IF(COUNTIF($A:$A,A{r})>1,"รหัสซ้ำ ","")&IF(LEN(A{r})>16,"ยาว "&LEN(A{r})&'
                       f'" ตัว (รายงานข้อความตัดเหลือ 16) ","")&IF(AND(G{r}="ชิ้นงานผลิต",Z{r}=""),'
                       f'"ยังไม่ผูกรหัสนับสต็อก",""))')
    style_body(ws, 2, n + 1, len(MASTER_HEADERS))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(MASTER_HEADERS))}{n + 1}"
    ws.conditional_formatting.add(f"AC2:AC{n + 1}", FormulaRule(formula=['LEN(AC2)>0'],
                                                                fill=PatternFill("solid", fgColor="FDECEC")))

    # ---- code groups -----------------------------------------------------------------
    fams = sorted({code_family(i.code) for i in express.items if i.code},
                  key=lambda f: (not f[:1].isdigit(), f))
    wf = sheet(FAMILIES, ["กลุ่มรหัส", "ความหมาย", "ประเภท", "จำนวนรายการ", "ผูกรหัสนับสต็อกแล้ว"],
               [12, 34, 14, 14, 18])
    for r, fam in enumerate(fams, start=2):
        wf.cell(r, 1, fam)
        wf.cell(r, 2, CODE_FAMILY_NAMES.get(fam, "(ยังไม่ระบุ)"))
        wf.cell(r, 2).fill = input_fill
        wf.cell(r, 3, item_kind(fam + "-"))
        wf.cell(r, 4, f"=COUNTIF({MASTER}!$E:$E,A{r})")
        wf.cell(r, 5, f'=COUNTIFS({MASTER}!$E:$E,A{r},{MASTER}!$Z:$Z,"<>")')
    style_body(wf, 2, len(fams) + 1, 5)

    # ---- counted stock (file 10) -----------------------------------------------------
    wc = sheet(COUNTED, ["รหัสนับสต็อก", "รายการ", "ยอดคงเหลือ", "หน่วย", "หมายเหตุ", "Minimum", "รหัส Express"],
               [16, 40, 11, 8, 18, 10, 22])
    for r, item in enumerate(counted.items, start=2):
        for c, v in enumerate([item.code, item.name, item.on_hand, item.unit, item.remark, item.min_qty or None],
                              start=1):
            if v not in (None, ""):
                wc.cell(r, c, v)
        wc.cell(r, 7, f'=IFERROR(INDEX({MASTER}!$A:$A,MATCH(A{r},{MASTER}!$Z:$Z,0)),"ยังไม่จับคู่")')
    style_body(wc, 2, len(counted.items) + 1, 7)

    # ---- model BOM codes vs Express ------------------------------------------------------
    b = SimpleNamespace(font=font, input_fill=input_fill, body_font=body_font, sheet=sheet, header=header,
                        style_body=style_body, master=MASTER, master_rows=n)
    model_tabs: list[str] = []
    if resolved is not None:
        from .partlist_express import write_model_tabs, write_review_tab

        before = set(wb.sheetnames)
        links = write_model_tabs(wb, b, models, resolved)
        model_tabs = [t for t in wb.sheetnames if t not in before]
        r = write_review_tab(wb, b, BOMCHECK, links) + 2
    else:
        r = _write_bom_check(sheet, style_body, models)
    balance_tabs = []
    if balances:
        from .partlist_express import write_balance_tab

        for name, stock in balances.items():
            title = f"{name} (Express)"[:31]
            write_balance_tab(wb, b, title, stock, express)
            balance_tabs.append(title)

    # ---- form for new items ---------------------------------------------------------------
    wn = sheet(FORM, FORM_HEADERS, [20, 34, 18, 7, 10, 9, 9, 8, 9, 8, 8, 10, 14, 20, 12, 11, 30], header_row=3)
    wn.cell(1, 1, "ฟอร์มขอเปิดรหัสสินค้าใหม่ใน Express (4.สินค้า → รายละเอียดสินค้า → เพิ่มใหม่)").font = \
        Font(name=font, bold=True, size=13)
    wn.cell(2, 1, "ช่องสีเหลือง = กรอก · แถวที่ 4 เป็นตัวอย่าง · ตรวจรหัส คำนวณให้อัตโนมัติ").font = \
        Font(name=font, italic=True, size=9, color="6B7A90")
    example = ["RRS-T-024-L-90", "บังอาย T90(ซ้าย)", "", "21", "RRS-T", "ชิ้น", "ชิ้น", 1, "ชิ้น", 1, "ST01", 0, "",
               "ตัวอย่าง", "แอม", date.today().strftime("%d/%m/%Y")]
    last_form_row = max(60, len(form_rows or []) + 40)
    for rr in range(4, last_form_row + 1):
        for c in range(1, len(FORM_HEADERS)):
            wn.cell(rr, c).fill = input_fill
            wn.cell(rr, c).font = body_font
        wn.cell(rr, 17, f'=IF(A{rr}="","",TRIM(IF(COUNTIF({MASTER}!$A:$A,A{rr})>0,"มีรหัสนี้แล้ว ","")'
                        f'&IF(LEN(A{rr})>20,"ยาวเกิน 20 ตัว ","")&IF(LEN(A{rr})>16,"ยาวเกิน 16 (รายงานตัด) ","")'
                        f'&IF(ISNUMBER(SEARCH(" ",TRIM(A{rr}))),"มีช่องว่าง ","")&IF(B{rr}="","ยังไม่ใส่ชื่อ ","")'
                        f'&IF(COUNTIF($A$5:$A${last_form_row},A{rr})>1,"รหัสซ้ำในฟอร์ม ","")))')
    for c, v in enumerate(example, start=1):
        wn.cell(4, c, v)
        wn.cell(4, c).font = Font(name=font, size=10, italic=True, color="6B7A90")
    for rr, values in enumerate(form_rows or [], start=5):
        for c, v in enumerate(values[:len(FORM_HEADERS) - 1], start=1):
            if v not in (None, ""):
                wn.cell(rr, c, v)

    # ---- lists for data validation ------------------------------------------------------
    wl = sheet(LISTS, ["หมวด", "ชื่อหมวด", "กลุ่มรหัส", "หน่วย", "รหัสบ/ช"], [8, 28, 12, 10, 10])
    cats = sorted({(i.category, i.extra.get("category_name", "")) for i in express.items if i.category})
    for rr, (cat, name) in enumerate(cats or [("21", "รายการเหล็กสำหรับผลิตดั้มพ์")], start=2):
        wl.cell(rr, 1, cat)
        wl.cell(rr, 2, name)
    for rr, fam in enumerate(fams, start=2):
        wl.cell(rr, 3, fam)
    units = sorted({i.unit for i in express.items if i.unit} | set(UNITS))
    for rr, u in enumerate(units, start=2):
        wl.cell(rr, 4, u)
    accounts = sorted({i.account_group for i in express.items if i.account_group} or {"ST01"})
    for rr, a in enumerate(accounts, start=2):
        wl.cell(rr, 5, a)

    def validation(col_letter: str, source_col: str, count: int):
        dv = DataValidation(type="list", formula1=f"={LISTS}!${source_col}$2:${source_col}${count + 1}",
                            allow_blank=True, showErrorMessage=False)
        wn.add_data_validation(dv)
        dv.add(f"{col_letter}4:{col_letter}{last_form_row}")

    validation("D", "A", max(len(cats), 1))
    validation("E", "C", len(fams))
    for col in ("F", "G", "I"):
        validation(col, "D", len(units))
    validation("K", "E", len(accounts))

    # ---- densities ----------------------------------------------------------------------
    wd = sheet(DENSITY, ["กลุ่มวัสดุ", "kg/dm³", "หมายถึง"], [12, 10, 34])
    for rr, row in enumerate(_density_rows(densities), start=2):
        for c, v in enumerate(row, start=1):
            wd.cell(rr, c, v)
        wd.cell(rr, 2).fill = input_fill
    wd.cell(7, 1, "Source: ค่ามาตรฐานทั่วไป (เหล็ก 7.85, สแตนเลส 7.93, อลูมิเนียม 2.70) – แก้ได้ในช่องสีเหลือง")

    # ---- summary ------------------------------------------------------------------------
    wsum = sheet(SUMMARY, ["รายการ", "จำนวน"], [42, 12])
    rows = [
        ("สินค้าทั้งหมดใน Express", f"=COUNTA({MASTER}!$A:$A)-1"),
        ("วัตถุดิบ", f'=COUNTIF({MASTER}!$G:$G,"วัตถุดิบ")'),
        ("ชิ้นงานผลิต (RRS)", f'=COUNTIF({MASTER}!$G:$G,"ชิ้นงานผลิต")'),
        ("ชิ้นงานที่ผูกรหัสนับสต็อกแล้ว", f'=COUNTIFS({MASTER}!$G:$G,"ชิ้นงานผลิต",{MASTER}!$Z:$Z,"<>")'),
        ("ชิ้นงานที่ยังไม่ผูกรหัสนับสต็อก", f'=COUNTIFS({MASTER}!$G:$G,"ชิ้นงานผลิต",{MASTER}!$Z:$Z,"")'),
        ("รายการในไฟล์นับสต็อก", f"=COUNTA({COUNTED}!$A:$A)-1"),
        ("รายการนับสต็อกที่ยังไม่จับคู่กับ Express", f'=COUNTIF({COUNTED}!$G:$G,"ยังไม่จับคู่")'),
    ]
    if resolved is None:
        rows += [
            ("รหัสในแบบกระบะ ตรงกับ Express", f'=COUNTIF({BOMCHECK}!$H:$H,"ตรงกับ Express")'),
            ("รหัสในแบบกระบะ มีรุ่นย่อย ต้องเลือกรหัสเต็ม", f'=COUNTIF({BOMCHECK}!$H:$H,"มีรุ่นย่อย – เลือกรหัสเต็ม")'),
            ("รหัสในแบบกระบะ ไม่พบใน Express", f'=COUNTIF({BOMCHECK}!$H:$H,"ไม่พบใน Express")'),
        ]
    else:
        from .express_codes import ST_AUTO, ST_CHECK, ST_EXACT, ST_NONE

        rows += [
            ("แบบกระบะ / เฟรม", len(model_tabs)),
            ("แถวชิ้นส่วนในแบบทั้งหมด (แยกซ้าย/ขวา)", f"=COUNTA({BOMCHECK}!$A:$A)-1"),
            (f"  {ST_EXACT} (รหัสในแบบมีใน Express)", f'=COUNTIF({BOMCHECK}!$I:$I,"{ST_EXACT}")'),
            (f"  {ST_AUTO} (จากชื่อ รุ่น ความสูงคอก ซ้าย/ขวา)", f'=COUNTIF({BOMCHECK}!$I:$I,"{ST_AUTO}")'),
            (f"  {ST_CHECK} – เลือกรหัสในช่องสีเหลือง", f'=COUNTIF({BOMCHECK}!$I:$I,"{ST_CHECK}")'),
            (f"  {ST_NONE} – ขอเปิดรหัสใหม่", f'=COUNTIF({BOMCHECK}!$I:$I,"{ST_NONE}")+'
                                              f'COUNTIF({BOMCHECK}!$I:$I,"ไม่มีรหัสในแบบ")'),
            ("  แถวที่ตอนนี้มีรหัสใน Express แล้ว", f'=COUNTIF({BOMCHECK}!$J:$J,"มีใน Express")'),
            ("รายการในฟอร์มเพิ่มสินค้าใหม่", f'=COUNTA({FORM}!$A$5:$A${last_form_row})'),
        ]
    rows.append(("รายการที่ ตรวจรหัส พบข้อสังเกต", f'=COUNTIF({MASTER}!$AC:$AC,"?*")'))
    for rr, (label, formula) in enumerate(rows, start=2):
        wsum.cell(rr, 1, label).font = body_font
        wsum.cell(rr, 2, formula).font = Font(name=font, size=10, bold=True)

    # ---- guide --------------------------------------------------------------------------
    guide.column_dimensions["A"].width = 110
    lines = [
        ("ฟอร์มทะเบียนสินค้า – จาก Express รายงานรายละเอียดสินค้า (STMASR0200)", True),
        (f"Source: Express รายงาน {source}, ดึงข้อมูลวันที่ {date.today().strftime('%d/%m/%Y')} "
         f"· สต็อกนับ: ไฟล์ '10. รายการ Stock ชิ้นส่วนกระบะดั้มพ์.นับสต๊อค' · แบบกระบะ: '7. Parts List All Model'", False),
        ("", False),
        (f"{MASTER}: สินค้าทุกรายการใน Express หนึ่งแถวต่อหนึ่งรหัส (หน่วยย่อย/ใหญ่/ซื้อ/ขาย + ตัวคูณ, รหัสบ/ช, ราคามาตรฐาน)", False),
        ("   คอลัมน์ ความหมายกลุ่มรหัส, ความหนาแน่น, น้ำหนัก/แผ่น, ยอดคงเหลือ, หมายเหตุนับสต็อก, ตรวจรหัส เป็นสูตร", False),
        ("   รหัสนับสต็อก (สีเหลือง) จับคู่จากชื่อสินค้าที่ตรงกันในไฟล์นับสต๊อค – แก้ได้เอง", False),
        (f"{SUMMARY}: ตัวเลขสรุปทั้งหมดเป็นสูตร", False),
        (f"{FAMILIES}: ความหมายของรหัสแต่ละกลุ่ม (แก้ช่องสีเหลืองได้)", False),
        (f"{COUNTED}: ยอดจากไฟล์นับสต๊อค พร้อมรหัส Express ที่จับคู่ได้", False),
        (f"{BOMCHECK}: รหัสในแบบกระบะ (รหัสฐาน เช่น RRS-T-024) เทียบกับรหัสรุ่นย่อยใน Express (RRS-T-024-L-84 …)"
         if resolved is None else f"{BOMCHECK}: ทุกแถวของทุกแบบ พร้อมสถานะการจับคู่ – กรองคอลัมน์ การจับคู่ = ควรตรวจ", False),
        (f"{FORM}: ฟอร์มขอเปิดรหัสใหม่ กรอกช่องสีเหลือง เลือกหมวด/กลุ่มรหัส/หน่วย/รหัสบ/ช จากรายการ", False),
        ("", False),
        ("ช่องสีเหลือง = ข้อมูลที่กรอกหรือแก้ไขได้ · ช่องอื่นมาจาก Express หรือเป็นสูตร", False),
        ("นำเข้าไฟล์นี้กลับใน RRS Part List Studio ได้ (หน้า สโตร์ → นำเข้าไฟล์สต็อก หรือ Google Sheets → ดึงสต็อก)", False),
    ]
    if resolved is not None:
        lines[0] = ("Parts List All Model (รหัส Express) – แบบกระบะทุกแบบ ใช้รหัสสินค้าเดียวกับ Express", True)
        lines[9:9] = [
            (f"แท็บแบบกระบะ ({len(model_tabs)} แท็บ): รูปแบบเดิม ลำดับ | รหัส | รายการ | ภาพประกอบ | ขนาด | จำนวนที่ใช้ | หมายเหตุ", False),
            ("   รหัส = รหัสเต็มใน Express (เช่น RRS-T-024-L-84) · แถว 1+1 แยกเป็น ซ้าย 1 แถว ขวา 1 แถว", False),
            ("   การจับคู่: รหัสตรง / จับคู่อัตโนมัติ = ใช้ได้ · ควรตรวจ = ดูตัวเลือกอื่น แล้วเลือกรหัสจากรายการในช่องสีเหลือง"
             " · ไม่พบใน Express = ดูแท็บ ฟอร์มเพิ่มสินค้าใหม่", False),
            ("   ชื่อใน Express, หน่วย, ยอดคงเหลือ, ราคามาตรฐาน, ตรวจรหัส เป็นสูตร – เปลี่ยนรหัสแล้วค่าเปลี่ยนตาม", False),
        ] + [(f"{t}: ยอดคงเหลือเดิม พร้อมรหัส Express ของแต่ละรายการ (จับคู่จากชื่อ)", False) for t in balance_tabs]
    guide_fill_row = next(i for i, (t, _) in enumerate(lines, start=1) if t.startswith("ช่องสีเหลือง"))
    for rr, (text, bold) in enumerate(lines, start=1):
        guide.cell(rr, 1, text).font = Font(name=font, bold=bold, size=13 if rr == 1 else 10)
        if rr == guide_fill_row:
            guide.cell(rr, 1).fill = input_fill

    order = [GUIDE, SUMMARY, *model_tabs, BOMCHECK, FORM, MASTER, FAMILIES, COUNTED, *balance_tabs, LISTS, DENSITY]
    wb._sheets = [wb[t] for t in order] + [ws for ws in wb.worksheets if ws.title not in order]
    wb.active = 0
    wb.save(path)
    return {"items": n, "linked": sum(1 for i in express.items if any(_key(c) in counted_by_code
                                                                       for c in xref.lookup(i.code))),
            "counted": len(counted.items), "bom_rows": r - 2, "model_tabs": len(model_tabs)}


def _write_bom_check(sheet, style_body, models: list[ModelBOM]) -> int:
    """Base codes of the model sheets against Express (when the models were not resolved)."""
    wbm = sheet(BOMCHECK, ["แบบกระบะ", "ลำดับ", "รหัสในแบบ", "รายการ", "จำนวนที่ใช้", "พบรหัสตรง", "จำนวนรุ่นย่อยใน Express",
                           "สถานะ"], [24, 7, 18, 30, 10, 11, 14, 26])
    r = 2
    for model in models:
        for row in model.rows:
            for code in row.codes or [""]:
                wbm.cell(r, 1, model.name)
                wbm.cell(r, 2, row.item_no)
                if code:
                    wbm.cell(r, 3, base_code(code))
                wbm.cell(r, 4, row.name)
                wbm.cell(r, 5, row.qty_text or row.qty)
                wbm.cell(r, 6, f'=IF(C{r}="","",COUNTIF({MASTER}!$A:$A,C{r}))')
                wbm.cell(r, 7, f'=IF(C{r}="","",COUNTIF({MASTER}!$A:$A,C{r}&"-*"))')
                wbm.cell(r, 8, f'=IF(C{r}="","ไม่มีรหัสในแบบ",IF(F{r}>0,"ตรงกับ Express",IF(G{r}>0,'
                               f'"มีรุ่นย่อย – เลือกรหัสเต็ม","ไม่พบใน Express")))')
                r += 1
    style_body(wbm, 2, r - 1, 8)
    return r


def item_form_to_stock(rows: list[list]) -> StockList:
    """Read the ทะเบียนสินค้า tab back (Excel or Google Sheet)."""
    stock = StockList.from_rows(rows, source="ฟอร์มทะเบียนสินค้า")
    header = next((r for r in rows[:5] if "รหัสสินค้า" in [str(c) for c in r]), [])
    col = {str(h): i for i, h in enumerate(header)}
    data = [r for r in rows[rows.index(header) + 1:]] if header else []
    for item, raw in zip(stock.items, [r for r in data if any(str(c).strip() for c in r[:2])]):
        for key, label in (("unit_big", "หน่วยใหญ่"), ("factor_big", "ตัวคูณใหญ่"), ("unit_buy", "หน่วยซื้อ"),
                           ("factor_buy", "ตัวคูณซื้อ"), ("unit_sell", "หน่วยขาย"), ("factor_sell", "ตัวคูณขาย"),
                           ("count_code", "รหัสนับสต็อก"), ("category_name", "ชื่อหมวด")):
            if label in col and col[label] < len(raw) and str(raw[col[label]]).strip() not in ("", "None"):
                item.extra[key] = raw[col[label]]
        if "รหัสบ/ช" in col and col["รหัสบ/ช"] < len(raw):
            item.account_group = str(raw[col["รหัสบ/ช"]] or "").strip()
    stock.has_quantities = False
    return stock


def counted_links(stock: StockList) -> list[tuple[str, str, str]]:
    """(Express code, stock-count code, name) pairs carried in a filled-in form."""
    return [(i.code, str(i.extra["count_code"]), i.name) for i in stock.items if i.extra.get("count_code")]

