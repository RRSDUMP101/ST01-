"""Express item master (STMASR0200), Express codes for the model BOMs, and the combined workbook."""

import openpyxl
import pytest
from conftest import SAMPLES

from partlist_studio.core.company_files import ModelBOM, ModelRow, read_company_workbook, split_codes
from partlist_studio.core.express import read_express_file
from partlist_studio.core.express_codes import ST_AUTO, ST_EXACT, ST_NONE, resolve_models, resolved_to_model, summary
from partlist_studio.core.item_form import build_item_form, item_form_to_stock
from partlist_studio.core.partlist_express import build_partlist_express, match_names, suggest_code
from partlist_studio.core.stock import StockItem, StockList, find_variants


@pytest.fixture(scope="module")
def express():
    stock, report = read_express_file(SAMPLES / "express_item_detail_STMASR0200.txt")
    assert report.title == "รายงานรายละเอียดสินค้า"
    return stock


def _row(no, code, name, qty_text="1"):
    qty = sum(float(x) for x in qty_text.split("+"))
    return ModelRow(item_no=no, raw_code=code, codes=split_codes(code), name=name, qty_text=qty_text, qty=qty)


@pytest.fixture
def models():
    return [
        ModelBOM("กระบะคอกสูง3ราว", "รายการชิ้นส่วนประกอบกระบะ Toyota (คอกสูง 3 ราว)", [
            _row("1", "RRS-A-033", "ตะขอเกี่ยวฝาท้าย", "6"),
            _row("2", "RRS-T-023", "เสาหลัง 80", "1+1"),
            _row("3", "RRS-019", "ฝาข้าง", "2"),
            _row("4", "RRS-A-127", "ปะกับ", "2"),
        ]),
        ModelBOM("กระบะคอกสูง4ราว champ", "รายการชิ้นส่วนประกอบกระบะ Toyota Champ (คอกสูง 4 ราว )", [
            _row("1", "RRS-T-024", "บังอาย 84", "1+1"),
        ]),
    ]


def test_item_detail_report(express):
    assert len(express.items) == 14
    item = next(i for i in express.items if i.code == "RRS-T-024-R-84-CHAMP")  # 20 characters, touches the name
    assert item.name == "บังอาย Champ 84(ขวา)"
    assert item.unit == "ชิ้น" and item.account_group == "ST01" and item.category == "21"
    assert item.extra["unit_buy"] == "ชิ้น" and item.extra["factor_sell"] == 1
    plate = express.items[0]
    assert plate.code == "01-B00 -0003" and plate.unit == "แผ่น" and plate.remark == "ฝาข้าง"
    assert [i.code for i in find_variants("RRS-T-024", express)] == [
        "RRS-T-024-L-40-CHAMP", "RRS-T-024-L-84", "RRS-T-024-R-40-CHAMP", "RRS-T-024-R-84-CHAMP"]


def test_resolve_models(express, models):
    resolved = resolve_models(models, express)
    rows = resolved["กระบะคอกสูง3ราว"]
    by_no = {}
    for r in rows:
        by_no.setdefault(r.item_no, []).append(r)
    assert by_no["1"][0].status == ST_EXACT and by_no["1"][0].code == "RRS-A-033"
    # 1+1 -> one line per side, with the stall height of the model
    assert [(r.side, r.code, r.qty) for r in by_no["2"]] == [("ซ้าย", "RRS-A-102-L-80", 1), ("ขวา", "RRS-A-102-R-80", 1)]
    assert by_no["3"][0].code == "RRS-A-019-188-5-2-30"  # ฝาข้าง, never ฝาท้าย
    assert by_no["4"][0].status == ST_NONE and not by_no["4"][0].code
    champ = resolved["กระบะคอกสูง4ราว champ"]
    assert [r.code for r in champ] == ["RRS-T-024-L-84", "RRS-T-024-R-84-CHAMP"]
    assert all(r.status == ST_AUTO for r in champ)
    assert sum(summary(resolved).values()) == 7

    model = resolved_to_model(models[0], rows)
    assert [r.code for r in model.rows][:3] == ["RRS-A-033", "RRS-A-102-L-80", "RRS-A-102-R-80"]
    assert model.rows[1].name == "เสาหลัง 80 (ซ้าย)"
    assert model.rows[-1].code == "RRS-A-127"  # not in Express: the code of the sheet is kept
    assert "รหัสเดิม RRS-T-023" in model.rows[1].remark


def test_suggest_code_and_name_match(express, models):
    resolved = resolve_models(models, express)
    missing = next(r for r in resolved["กระบะคอกสูง3ราว"] if r.status == ST_NONE)
    assert suggest_code(models[0], missing, {"RRSA127"}) == "RRS-A-127-2"
    items = [StockItem(name="เหล็กแผ่นดำ 2.5 มม. 4x8 ฟุต"), StockItem(name="บังอาย 40"), StockItem(name="ไม่มีของนี้")]
    codes = [m[0] for m in match_names(items, express)]
    assert codes == ["01-B00 -0003", "RRS-A-024", ""]


def test_partlist_express_workbook(express, models, tmp_path):
    counted = StockList(items=[StockItem(code="01-0003-0099", name="เสาหลัง(ซ้าย) 80", on_hand=4, unit="ต้น")])
    steel = StockList(items=[StockItem(name="เหล็กแผ่นดำ 2.5 มม. 4x8 ฟุต", on_hand=15, unit="แผ่น",
                                       remark="ยกมา 15 · รับ 0 · จ่าย 0")])
    path = tmp_path / "Parts List All Model (รหัส Express).xlsx"
    stats = build_partlist_express(path, models, express, counted=counted, balances={"สต๊อกเหล็ก": steel})
    assert stats["model_tabs"] == 2 and stats["new_items"] == 1

    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[:5] == ["คำแนะนำ", "สรุป", "กระบะคอกสูง3ราว", "กระบะคอกสูง4ราว champ", "เทียบรหัสแบบกระบะ"]
    assert "สต๊อกเหล็ก (Express)" in wb.sheetnames
    ws = wb["กระบะคอกสูง3ราว"]
    assert [c.value for c in ws[3]][:7] == ["ลำดับ", "รหัส", "รายการ", "ภาพประกอบ", "ขนาด", "จำนวนที่ใช้", "หมายเหตุ"]
    assert ws["B5"].value == "RRS-A-102-L-80" and ws["H5"].value.startswith("=IF(B5")
    form = wb["ฟอร์มเพิ่มสินค้าใหม่"]
    assert form["A5"].value == "RRS-A-127" and form["B5"].value == "ปะกับ"
    assert wb["สต๊อกเหล็ก (Express)"]["A2"].value == "01-B00 -0003"

    # the program reads the model tabs back with the Express codes
    book = read_company_workbook(path)
    back = {m.name: m for m in book.models}
    assert set(back) == {"กระบะคอกสูง3ราว", "กระบะคอกสูง4ราว champ"}
    assert [r.code for r in back["กระบะคอกสูง4ราว champ"].rows] == ["RRS-T-024-L-84", "RRS-T-024-R-84-CHAMP"]
    assert back["กระบะคอกสูง3ราว"].title.startswith("รายการชิ้นส่วนประกอบกระบะ Toyota")


def test_item_form_round_trip(express, tmp_path):
    path = tmp_path / "form.xlsx"
    stats = build_item_form(express, path)
    assert stats["items"] == 14
    ws = openpyxl.load_workbook(path)["ทะเบียนสินค้า"]
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    rows[1][25] = "01-0001-0001"  # รหัสนับสต็อก typed by a person
    back = item_form_to_stock(rows)
    assert [i.code for i in back.items] == [i.code for i in express.items]
    assert back.items[0].extra["count_code"] == "01-0001-0001"
    assert back.items[1].extra["unit_buy"] == "ชิ้น"
