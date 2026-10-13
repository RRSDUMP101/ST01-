"""Reading and writing of plain tables (CSV / TXT / XLSX / XLS)."""

from __future__ import annotations

import csv
import io
from pathlib import Path

TABLE_SUFFIXES = {".csv", ".txt", ".tsv", ".xlsx", ".xlsm", ".xls"}

# Thai Windows exports from ZWCAD/Excel are often cp874 (TIS-620) rather than UTF-8.
_ENCODINGS = ["utf-8-sig", "cp874", "cp1252"]


def decode_bytes(raw: bytes) -> str:
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    for enc in _ENCODINGS:
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _sniff_delimiter(text: str) -> str:
    sample = "\n".join(text.splitlines()[:20])
    counts = {d: sample.count(d) for d in ("\t", ",", ";", "|")}
    return max(counts, key=counts.get) if any(counts.values()) else ","


def read_text_table(text: str) -> list[list[str]]:
    reader = csv.reader(io.StringIO(text), delimiter=_sniff_delimiter(text))
    return [[c.strip() for c in row] for row in reader]


def read_table(path: str | Path, sheet: str | None = None) -> list[list]:
    """Return all rows of the first (or named) sheet as lists of cell values."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
    elif suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        ws = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
        rows = [ws.row_values(i) for i in range(ws.nrows)]
    else:
        rows = read_text_table(decode_bytes(path.read_bytes()))
    return _trim(rows)


def _trim(rows: list[list]) -> list[list]:
    """Drop empty rows and trailing empty cells."""
    out = []
    for row in rows:
        cells = list(row)
        while cells and (cells[-1] is None or str(cells[-1]).strip() == ""):
            cells.pop()
        if cells:
            out.append(cells)
    return out


def write_table(path: str | Path, header: list[str], rows: list[list], title_rows: list[list] | None = None,
                sheet_title: str = "Sheet1") -> None:
    """Write a table to .csv (UTF-8 with BOM, opens correctly in Excel/ZWCAD) or .xlsx."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        write_workbook(path, [Sheet(sheet_title, header, rows, title_rows or [])])
        return
    delimiter = "\t" if path.suffix.lower() in (".txt", ".tsv") else ","
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=delimiter)
        writer.writerows(title_rows or [])
        writer.writerow(header)
        writer.writerows(rows)


class Sheet:
    def __init__(self, title: str, header: list[str], rows: list[list], title_rows: list[list] | None = None,
                 images: dict[int, str] | None = None, image_col: int | None = None, blank_after_title: bool = True):
        self.title, self.header, self.rows = title, header, rows
        self.title_rows = title_rows or []
        self.images = images or {}  # data row index -> image file
        self.image_col = image_col
        self.blank_after_title = blank_after_title


def write_workbook(path: str | Path, sheets: list[Sheet]) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(sheet.title[:31])
        for row in sheet.title_rows:
            ws.append(row)
        if sheet.title_rows:
            ws.cell(1, 1).font = Font(bold=True, size=14)
            if sheet.blank_after_title:
                ws.append([])
        ws.append(sheet.header)
        header_row = ws.max_row
        for cell in ws[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.rows:
            ws.append(row)
        ws.freeze_panes = ws.cell(header_row + 1, 1)
        for i, col in enumerate(ws.iter_cols(min_row=header_row, max_row=ws.max_row), start=1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 45)
        if sheet.images and sheet.image_col is not None:
            _add_images(ws, sheet, header_row)
    wb.save(path)


def _add_images(ws, sheet: Sheet, header_row: int) -> None:
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import get_column_letter

    letter = get_column_letter(sheet.image_col + 1)
    ws.column_dimensions[letter].width = 24
    for idx, file in sheet.images.items():
        if not file or not Path(file).is_file():
            continue
        try:
            img = XLImage(file)
        except Exception:  # noqa: BLE001 - unreadable picture, keep the row
            continue
        scale = min(160 / max(img.width, 1), 90 / max(img.height, 1), 1)
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        row = header_row + 1 + idx
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 15, img.height * 0.78)
        ws.add_image(img, f"{letter}{row}")
