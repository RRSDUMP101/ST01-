"""Catalog of the store: item types, assembly groups, the six numbers per part, readiness, tickets and labels."""

import openpyxl
import pytest
from conftest import SAMPLES

from partlist_studio.core import catalog, dwg_io, laser, nesting, production, qr
from partlist_studio.core.express import read_express_file
from partlist_studio.core.models import Part
from partlist_studio.core.stock import StockItem, StockList


@pytest.mark.parametrize("name, group", [
    ("เฟรมล่าง TOYOTA (New)", "FL"),
    ("เหล็กซับคัดซี (230 ลงมา)", "FL"),
    ("เฟรมล่าง TOYOTA -Revo(มีรางรถไฟ)", "FL"),  # a frame is a frame
    ("หน้าเฟรมล่าง Champ", "FL"),
    ("หน้าเฟรมบน T-champ", "FU"),
    ("รองท้ายเฟรมบน T-รางรถไฟ", "FU"),
    ("กระดูกข้าง Champ(ซ้าย)", "FU"),  # the ribs are counted with the upper frame (01-0004)
    ("กระดูกกลางพื้นเรียบเบอร์ 4", "FU"),
    ("บังอาย Champ 84(ซ้าย)", "FRONT"),
    ("เสาโรบาร์พื้นเรียบ", "FRONT"),
    ("เสาหน้าพื้นเรียบ(ขวา)", "POST"),  # "พื้นเรียบ" is also a bed model: the part type wins
    ("เสาปิดฝาท้ายคอกสูง(ซ้าย)", "POST"),
    ("ฝาข้าง(209)กระบะ2.50", "SIDE"),
    ("เหล็กฉากราวล่าง(ซ้าย)", "SIDE"),
    ("หูล็อกฝาท้าย ซ้าย", "TAIL"),  # where it is assembled
    ("เหล็กฉากไฟท้ายพื้นเรียบ", "TAIL"),
    ("พื้นเรียบ(เครน)T-Revo", "FLOOR"),
    ("ซองรางรถไฟ(ยาว)", "RAIL"),
    ("ตัวดึงรางรถไฟ", "RAIL"),
    ("หูช้าง(ขวา)", "FIX"),
    ("บู๊ชหูบน", "FIX"),
    ("กันระเบิด", "ACC"),
    ("ปะกับ", ""),
])
def test_assembly_group_from_the_name(name, group):
    assert catalog.group_by_name(name) == group


def test_item_types_store_codes_and_overrides():
    assert catalog.item_type("RRS-T-024-L-84", "บังอาย", "ชิ้น") == "PART"
    assert catalog.item_type("RRS-A-002", "หน้าโปร", "ชุด") == "SUB"
    assert catalog.item_type("01-B00 -0004", "เหล็กแผ่นดำ 3.2 มม.4x8ฟุต") == "RM"
    assert catalog.item_type("60-001", "ลวดเชื่อม 1.2") == "CON"
    assert catalog.item_type("20-001", "น็อตหัวเหลี่ยม M10") == "BUY"
    assert catalog.item_type("FG-HB3-T-230", "กระบะคอกสูง") == "FG"
    assert catalog.assembly_group("01-0009-0012", "ชิ้นใหม่") == "FIX"  # the stock-count group of the store
    assert catalog.assembly_group("01-0011-0001", "ชิ้นใหม่") == ""  # 0011 พื้นเรียบ is a bed model, not a group
    cat = catalog.Catalog()
    cat.set_group("rrs-t-024-l-84", "POST")
    assert cat.group_of("RRS-T-024-L-84", "บังอาย") == "POST"
    cat.set_group("RRS-T-024-L-84", None)  # back to the rules
    assert cat.group_of("RRS-T-024-L-84", "บังอาย") == "FRONT"
    cat.set_location("RRS-T-024-L-84", " P-FRONT-02 ")
    item = StockItem(code="RRS-A-033", location="ชั้น 3")
    assert cat.location_of("RRS T 024 L 84") == "P-FRONT-02" and cat.location_of("RRS-A-033", item) == "ชั้น 3"
    back = catalog.Catalog.from_dict(cat.to_dict())
    assert back.locations == cat.locations
    cat.set_location("RRS-T-024-L-84", "")
    assert not cat.locations
    assert catalog.brand_of("RRS-T-024-L-84") == "Toyota" and catalog.side_of("RRS-T-024-L-84", "") == "ซ้าย"


def test_every_made_part_of_the_express_sample_has_a_group():
    stock, _ = read_express_file(SAMPLES / "express_item_detail_STMASR0200.txt")
    made = [i for i in stock.items if catalog.item_type(i.code, i.name, i.unit_name) in catalog.MADE_TYPES]
    assert made and all(catalog.assembly_group(i.code, i.name) for i in made)


def _jobs():
    a = production.job_from_parts("SO-A", [], 3, None, False)
    a.parts = [production.JobPart("1", "RRS-T-024-L-84", "บังอาย Champ 84(ซ้าย)", 2, 6, from_stock=2,
                                  route=["ตัด", "พับ"], done={"ตัด": 3, "พับ": 1}),
               production.JobPart("2", "RRS-X-001", "ขาจับ", 1, 3, route=["ตัด"])]
    b = production.job_from_parts("SO-B", [], 2, None, False)
    b.parts = [production.JobPart("1", "RRS-A-102-L-80", "เสาหลัง(ซ้าย) 80", 2, 4, from_stock=4, route=[])]
    closed = production.job_from_parts("SO-OLD", [], 1, None, False)
    closed.status = production.JOB_DELIVERED
    closed.parts = [production.JobPart("1", "RRS-A-102-L-80", "เสาหลัง(ซ้าย) 80", 9, 9)]
    return [a, b, closed]


def test_six_numbers_of_every_part():
    stock = StockList(items=[
        StockItem(code="RRS-T-024-L-84", name="บังอาย Champ 84(ซ้าย)", unit="ชิ้น", on_hand=4, min_qty=2),
        StockItem(code="RRS-A-102-L-80", name="เสาหลัง(ซ้าย) 80", unit="ต้น", on_hand=1),
        StockItem(code="RRS-A-033", name="ตะขอเกี่ยวฝาท้าย", unit="ชิ้น", on_hand=12),
        StockItem(code="01-B00-0004", name="เหล็กแผ่นดำ 3.2 มม.", unit="แผ่น", on_hand=5)])  # raw: not listed
    cat = catalog.Catalog()
    cat.set_location("RRS-A-033", "P-TAIL-01")
    rows = {r.code: r for r in catalog.part_stock(stock, _jobs(), cat)}
    assert set(rows) == {"RRS-T-024-L-84", "RRS-A-102-L-80", "RRS-A-033", "RRS-X-001"}
    eye = rows["RRS-T-024-L-84"]  # 6 for SO-A: 2 from stock + 4 made (3 cut, 1 bent)
    assert (eye.on_hand, eye.need, eye.incoming, eye.reserved, eye.free, eye.short) == (4, 6, 4, 2, 2, 0)
    assert eye.wip == {"ตัด": 1, "พับ": 2} and eye.wip_text == "รอตัด 1 · รอพับ 2" and eye.status == "ระหว่างผลิต"
    post = rows["RRS-A-102-L-80"]  # SO-B takes 4 from stock, only 1 on the shelf; the delivered job is not counted
    assert (post.need, post.incoming, post.short, post.status, post.jobs) == (4, 0, 3, "ขาด", ["SO-B"])
    hook = rows["RRS-A-033"]
    assert hook.group == "TAIL" and hook.location == "P-TAIL-01" and hook.free == 12 and hook.status == "พอ"
    assert rows["RRS-X-001"].item is None and rows["RRS-X-001"].group == ""
    ordered = catalog.part_stock(stock, _jobs(), cat)
    assert [r.group for r in ordered] == ["FRONT", "POST", "TAIL", ""]  # assembly order, no group last
    assert catalog.part_stock_rows(ordered)[0][:3] == ["หน้าโปร / บังอาย / โรลบาร์", "RRS-T-024-L-84",
                                                         "บังอาย Champ 84(ซ้าย)"]


def test_readiness_kitting_and_files(tmp_path):
    job = _jobs()[0]
    job.record(job.parts[1], production.DONE, 0)  # ขาจับ made
    stock = StockList(items=[StockItem(code="RRS-T-024-L-84", name="บังอาย", location="P-FRONT-01")])
    groups = catalog.job_readiness(job, None, stock)
    eye, other = groups
    assert eye.group == "FRONT" and (eye.required, eye.ready, eye.complete) == (6, 3, False)  # 2 from stock + 1 made
    assert "×3 (รอตัด 1 · รอพับ 2)" in eye.missing_text() and eye.lines[0].location == "P-FRONT-01"
    assert other.group == "" and other.complete and other.lines[0].where == "เสร็จ รอรับเข้า"
    assert catalog.job_ready_pct(groups) == pytest.approx(6 / 9)
    rows = catalog.readiness_rows(groups)
    assert rows[0][3] == "50%" and rows[0][4] == "ขาด" and rows[1][4] == "พร้อม"
    catalog.export_kitting(job, groups, tmp_path / "kit.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "kit.xlsx")
    assert wb.sheetnames == ["ใบจัดชุด", "สรุปความพร้อม"]
    values = [c for row in wb["ใบจัดชุด"].iter_rows(values_only=True) for c in row]
    assert "RRS-T-024-L-84" in values and "P-FRONT-01" in values
    catalog.export_part_stock(catalog.part_stock(stock, [job]), tmp_path / "parts.xlsx")
    assert openpyxl.load_workbook(tmp_path / "parts.xlsx").sheetnames == ["ชิ้นส่วนรอประกอบ", "สรุปตามกลุ่ม"]


def test_tickets_labels_and_qr():
    data = dwg_io.read_drawing(SAMPLES / "RRS-T-SET_frame_parts.dxf")
    parts = [c.part for c in dwg_io.drawing_to_candidates(data)]
    lib = laser.FlatLibrary()
    lib.add_parts(parts)
    job = production.job_from_parts("SO6910-015", parts, 3, None, False, title="คอกสูง 3 ราว", due="2026-10-30")
    job.parts[1].from_stock = job.parts[1].required  # from stock: no ticket
    queue = laser.CutQueue()
    demand = laser.demand_from_jobs([job], lib)
    laser.make_programs(nesting.plan(demand.parts, None, laser.LaserParams().nest_params()), queue, laser.LaserParams())
    cat = catalog.Catalog()
    cat.set_location("RRS-T-001", "P-FU-01")
    tickets = catalog.job_tickets(job, None, cat, None, lib, queue.programs)
    assert [t.code for t in tickets] == ["RRS-T-001", "RRS-T-003", "RRS-T-004", "RRS-T-005"]
    t = tickets[0]
    assert t.no == "SO6910-015/01" and t.qty == 48 and t.due == "30/10/2569" and (t.material, t.thickness) == ("Steel", 3)
    assert t.to == "เก็บที่ P-FU-01 · กลุ่มเฟรมบน" and t.outline and t.program.startswith("LC")
    assert t.route == [("ตัด", 0, 48), ("พับ", 0, 48)] and "พื้นที่สุทธิ" not in t.remark
    scan = qr.parse_payload(t.payload)
    assert (scan.kind, scan.job_no, scan.line, scan.code, scan.qty) == ("job", "SO6910-015", 1, "RRS-T-001", 48)
    # the form: DWG code, Express code, 2D of the drawing with its size, vehicle, body model; the Part List
    # details (material, size) in the QR after the six fields the program reads
    assert (t.drawing_code, t.express_code, t.express_note) == ("RRS-T-001", "", "")  # no stock file given
    assert t.size == "150 × 60 มม." and t.vehicle == "Toyota" and t.body == "คอกสูง 3 ราว"
    assert scan.details == ["RRS-T-001", "", "Steel t3", "150x60", "เฟรมบน"]
    stock = StockList(items=[StockItem(code="RRS-A-001-250", name="เฟรมบน 250", on_hand=1)])
    job.brand, job.vehicle = "T", "REVO"
    [t2] = catalog.job_tickets(job, [0], cat, stock, lib, queue.programs)
    assert (t2.express_code, t2.vehicle) == ("RRS-A-001-250", "Toyota · REVO")  # 001, A = every brand
    assert qr.parse_payload(t2.payload).details[:2] == ["RRS-T-001", "RRS-A-001-250"]
    [t3] = catalog.job_tickets(job, [2], cat, stock, lib, queue.programs)
    assert (t3.express_code, t3.express_note) == ("", "ยังไม่มีในสต็อก")
    from partlist_studio.ui.printing import TICKET_ROWS, ticket_rows

    rows = ticket_rows(t3)
    assert [label for label, _v in rows][:7] == TICKET_ROWS and rows[1][1] == "ยังไม่มีในสต็อก"
    # M = วัสดุ, then t = ความหนาเหล็ก (มม.) on the next line; the size line has no thickness
    labels = [label for label, _v in rows]
    assert labels.index("t = ความหนาเหล็ก") == labels.index("M = วัสดุ") + 1
    assert dict(rows)["ขนาด"] == "100 × 80 มม." and dict(rows)["M = วัสดุ"].startswith("SS400")
    assert "t4" not in dict(rows)["M = วัสดุ"] and dict(rows)["t = ความหนาเหล็ก"] == "4 มม."
    matrix = qr.qr_matrix(t.payload)
    assert len(matrix) >= 21 and all(len(row) == len(matrix) for row in matrix)

    [label] = catalog.shelf_labels([r for r in catalog.part_stock(StockList(), [job], cat) if r.code == "RRS-T-001"])
    assert label.location == "P-FU-01" and label.group == "เฟรมบน" and label.brand == "Toyota"
    assert qr.parse_payload(label.payload).code == "RRS-T-001" and qr.parse_payload(label.payload).kind == "item"
    assert qr.parse_payload("RRS-T-001").kind == "text"
    assert qr.ticket_payload("SO|1", 2, "A|B", 1.5) == "RRS1|J|SO/1|2|A/B|1.5"  # the separator never comes from data


def test_job_keeps_the_dwg_code_next_to_the_stock_code():
    """A Part List row tied to an Express code, and a model row: the ticket shows both codes."""
    parts = [Part(part_no="RRS-T-039", name="เต้า1", stock_code="RRS-A-039-NO1-PR", qty=2, spec="t5",
                  material="SS400", thickness=5),
             Part(part_no="RRS-T-024-L-84", name="บังอาย Champ 84(ซ้าย)", stock_code="RRS-T-024-L-84",
                  remark="รหัสเดิม RRS-T-024")]
    job = production.job_from_parts("SO-1", parts, 1, None, False, body="กระบะดั้มพ์คอกสูง 3 ราว", brand="T")
    first, model_row = job.parts
    assert (first.code, first.drawing_code, first.spec, first.material) == ("RRS-A-039-NO1-PR", "RRS-T-039", "t5", "SS400")
    assert (model_row.code, model_row.drawing_code) == ("RRS-T-024-L-84", "RRS-T-024")
    again = production.Job.from_dict(job.to_dict())
    assert again.body == "กระบะดั้มพ์คอกสูง 3 ราว" and again.parts[0].drawing_code == "RRS-T-039"
    t, t2 = catalog.job_tickets(again)
    assert (t.drawing_code, t.express_code, t.material, t.thickness, t.size) == ("RRS-T-039", "RRS-A-039-NO1-PR",
                                                                       "SS400", 5, "t5")
    assert (t.vehicle, t.body) == ("Toyota", "กระบะดั้มพ์คอกสูง 3 ราว")
    assert (t2.drawing_code, t2.express_code) == ("RRS-T-024", "RRS-T-024-L-84")
